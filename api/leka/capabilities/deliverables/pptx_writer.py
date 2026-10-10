"""Slides, written as an editable PowerPoint: native shapes, connectors and
tables (the template asks for them, never a picture), Arial, drawn in the
slide's order — the first item at the back, so a lifeline passes behind
the boxes.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Length, Pt

from leka.capabilities.deliverables.scene import (
    HEIGHT,
    WIDTH,
    Box,
    Slide,
    Stroke,
    Table,
    side_inset,
)

FONT = "Arial"

_SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "hexagon": MSO_SHAPE.HEXAGON,
    "chevron": MSO_SHAPE.CHEVRON,
    "can": MSO_SHAPE.CAN,
}
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _in(value: float) -> Length:
    return Inches(value)


def _rgb(hex_colour: str) -> RGBColor:
    red, green, blue = bytes.fromhex(hex_colour)
    return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]  # untyped in python-pptx


def _element(tag: str, **attributes: str) -> Any:
    element = OxmlElement(tag)
    for name, value in attributes.items():
        element.set(name, value)
    return element


def _unstyled(shape: Any) -> None:
    """No theme style: no shadow, no theme colour behind ours."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _write_text(
    frame: Any, text: str, size: float, bold: bool, italic: bool, colour: str, align: str
) -> None:
    frame.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = _ALIGN[align]
        run = paragraph.add_run()
        run.text = line
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.italic = italic
        font.color.rgb = _rgb(colour)


def _box(slide: Any, box: Box) -> None:
    shape = slide.shapes.add_shape(
        _SHAPES[box.shape], _in(box.x), _in(box.y), _in(box.w), _in(box.h)
    )
    _unstyled(shape)
    if box.shape == "round":
        shape.adjustments[0] = 0.16667
    if box.fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(box.fill)
    else:
        shape.fill.background()
    if box.line:
        shape.line.color.rgb = _rgb(box.line)
        shape.line.width = Pt(box.line_width)
        if box.dashed:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        shape.line.fill.background()
    frame = shape.text_frame
    inset = _in(box.inset)
    frame.margin_top = frame.margin_bottom = inset
    frame.margin_left = frame.margin_right = _in(side_inset(box.shape, box.w, box.h, box.inset))
    frame.vertical_anchor = _ANCHOR[box.valign]
    frame.auto_size = None
    if box.text:
        _write_text(frame, box.text, box.size, box.bold, box.italic, box.color, box.align)


def _stroke(slide: Any, stroke: Stroke) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, _in(stroke.x1), _in(stroke.y1), _in(stroke.x2), _in(stroke.y2)
    )
    _unstyled(line)
    line.line.color.rgb = _rgb(stroke.color)
    line.line.width = Pt(stroke.width)
    if stroke.dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if stroke.arrow:
        ln = line.line._get_or_add_ln()
        ln.append(_element("a:tailEnd", type="triangle", w="med", len="med"))


def _borders(cell: Any, colour: str = "BFBFBF") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for i, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = _element(side, w="6350")
        fill = _element("a:solidFill")
        fill.append(_element("a:srgbClr", val=colour))
        ln.append(fill)
        tc_pr.insert(i, ln)


def _table(slide: Any, table: Table) -> None:
    rows, cols = 1 + len(table.rows), len(table.widths)
    frame = slide.shapes.add_table(
        rows, cols, _in(table.x), _in(table.y), _in(sum(table.widths)), _in(sum(table.heights))
    )
    grid = frame.table
    grid.first_row = True
    grid.horz_banding = False
    for c, width in enumerate(table.widths):
        grid.columns[c].width = _in(width)
    for r, height in enumerate(table.heights):
        grid.rows[r].height = _in(height)

    def fill(cell: Any, text: str, fill: str, colour: str, bold: bool) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
        _borders(cell)
        margin = _in(table.margin)
        cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = margin
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text_frame.text = ""
        _write_text(cell.text_frame, text, table.size, bold, False, colour, "left")

    for c, text in enumerate(table.header):
        fill(grid.cell(0, c), text, table.header_fill, table.header_color, True)
    for r, row in enumerate(table.rows, 1):
        ground = table.zebra if r % 2 == 0 else "FFFFFF"
        for c, cell in enumerate(row):
            fill(grid.cell(r, c), cell.text, cell.fill or ground, cell.color, cell.bold)


def write_pptx(slides: list[Slide]) -> bytes:
    deck = Presentation()
    deck.slide_width = _in(WIDTH)
    deck.slide_height = _in(HEIGHT)
    blank = deck.slide_layouts[6]
    for scene in slides:
        page = deck.slides.add_slide(blank)
        for item in scene.items:
            if isinstance(item, Box):
                _box(page, item)
            elif isinstance(item, Stroke):
                _stroke(page, item)
            else:
                _table(page, item)
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


__all__ = ["write_pptx"]
