"""The architect's deck, drawn (capabilities/deliverables) — pure: no
database, no model.

What these protect, as the agent's template asks (SKILL.md, v18):
- The slides in their imperative order — title and contents, the journey
  in plain words, the rules, the application view when asked, one slide per
  sub-process, the RACI, the points to confirm, the sources LAST — titled as
  the template titles them; the contents list the sections, the processes
  on one entry, the application view its own when there is one.
- Nothing leaves the slide; what the sources lack is orange; a process with
  more lines than a slide holds is split (1/2), (2/2), its business layer
  repeated; the RACI fills its slide before the next.
- On a process slide the lifelines are behind the bands' titles and the
  service boxes; a red frame and a red line lead to the front.
- The PowerPoint is native — shapes, connectors, tables — and opens; the
  PDF has a page per slide; both carry the same text.
"""

from __future__ import annotations

import json
from io import BytesIO

from pptx import Presentation

from leka.capabilities.deliverables import ArchitectureDeck, render
from leka.capabilities.deliverables.deck import (
    AppBox,
    AppView,
    Band,
    BusinessLayer,
    Call,
    Column,
    Contact,
    Control,
    Flow,
    Grouping,
    Line,
    Overview,
    ProcessView,
    Rule,
    SourceRow,
    Step,
)
from leka.capabilities.deliverables.layout import deck_slides
from leka.capabilities.deliverables.pdf_writer import plain
from leka.capabilities.deliverables.scene import (
    HEIGHT,
    LIFELINE,
    ORANGE,
    RED,
    SERVICE,
    WIDTH,
    Box,
    Slide,
    Stroke,
    Table,
)

COLUMNS = [
    Column(name="API PORTAIL BACKEND", subtitle="portail_ws"),
    Column(name="NEXUS", subtitle="middleware d'échanges"),
    Column(name="FACTURATION", subtitle="ERP"),
    Column(name="ORDONNANCEUR", missing=True),
]


def chain(*calls: tuple[str, str]) -> Line:
    return Line(calls=[Call(column=c, label=label) for c, label in calls])


def deck(*, lines: int = 2, app_view: bool = True, contacts: int = 6) -> ArchitectureDeck:
    process = [
        ProcessView(
            letter=letter,
            label=label,
            front="DIGITAL ESPACE CLIENT",
            front_label="FRONT WEB",
            channel_label="Espace Client",
            columns=COLUMNS,
            bands=[
                Band(
                    title="Vérification du solde",
                    lines=[
                        chain(
                            ("API PORTAIL BACKEND", "ws/paiement/carte"),
                            ("NEXUS", "ws_lireSolde (solde)"),
                            ("FACTURATION", "BAPI"),
                        )
                    ]
                    * lines,
                ),
                Band(
                    title="Ordonnancement",
                    lines=[
                        Line(calls=[Call(column="ORDONNANCEUR", label="planifie", missing=True)])
                    ],
                ),
            ],
        )
        for letter, label in (("A", "Initialisation du paiement"), ("B", "Paiement"))
    ]
    return ArchitectureDeck(
        parcours="Paiement CB",
        perimetre="Périmètre Digital - Espace Client",
        business=BusinessLayer(
            actor="Client",
            trigger="Choisit de payer par Carte Bancaire",
            process_name="Paiement par CB",
            channel="Espace Client",
            groupings=[
                Grouping(
                    letter="A",
                    label="Initialisation du paiement",
                    steps=[Step(label="Affiche le solde"), Step(label="Choisit le montant")],
                ),
                Grouping(
                    letter="B",
                    label="Paiement",
                    steps=[
                        Step(label="Saisit la carte"),
                        Step(label="Paiement accepté", kind="event"),
                        Step(label="Paiement refusé", branch="below"),
                        Step(label="Notifie le client", missing=True),
                    ],
                ),
            ],
        ),
        overview=Overview(objectif="Payer sa facture par carte.", points_attention="Relance ?"),
        rules=[Rule(etape="A - Solde", regle="Le bouton Payer suit le solde.", statut="Documenté")],
        app_view=AppView(
            apps=[
                AppBox(name="Espace Client", kind="perimetre"),
                AppBox(name="NEXUS", kind="middleware"),
                AppBox(name="FACTURATION", kind="connexe", stores=["Factures"]),
                AppBox(name="Prestataire de paiement", kind="externe"),
            ],
            flows=[
                Flow(n="1", source="Espace Client", target="NEXUS", interface="API REST"),
                Flow(n="2", source="NEXUS", target="FACTURATION", interface="BAPI", data="solde"),
                Flow(n="3", source="NEXUS", target="Prestataire de paiement", status="À confirmer"),
            ],
        )
        if app_view
        else None,
        processes=process,
        raci=[
            Contact(application="NEXUS", role="Conduite applicative", raci="R")
            for _ in range(contacts)
        ],
        raci_notes="L'annuaire n'a pas répondu pour NEXUS.",
        controls=[Control(manque="ORDONNANCEUR ?", localisation="Process A - bande 2")],
        sources=[SourceRow(source="Confluence MO", elements="Parcours", statut="Consulté")],
    )


def texts(slide: Slide) -> list[str]:
    found: list[str] = []
    for item in slide.items:
        if isinstance(item, Box) and item.text:
            found.append(item.text)
        elif isinstance(item, Table):
            found.extend(item.header)
            found.extend(cell.text for row in item.rows for cell in row)
    return found


def title(slide: Slide) -> str:
    return next(t for t in texts(slide))


# ── the template ────────────────────────────────────────────────────────────


def test_the_slides_follow_the_templates_order_and_titles() -> None:
    slides = deck_slides(deck())
    assert [title(s) for s in slides] == [
        "Architecture applicative - Parcours Paiement CB",
        "Le parcours en clair - Paiement CB",
        "Règles de gestion - Paiement CB",
        "Vue applicative - Paiement CB",
        "Parcours Paiement CB - Vue globale - process A : A - INITIALISATION DU PAIEMENT",
        "Parcours Paiement CB - Vue globale - process B : B - PAIEMENT",
        "Contacts et responsabilités (RACI) - parcours Paiement CB / périmètre Digital",
        "Points de contrôle / éléments à confirmer (éléments ORANGE des schémas)",
        "Sources consultées et limites",
    ]
    contents = texts(slides[0])
    assert "Sommaire" in contents
    assert "4 - Les process : A - Initialisation du paiement, B - Paiement" in contents
    assert "3 - Vue applicative - Paiement CB" in contents
    assert contents[-1] == "7 - Sources consultées et limites"
    # Its subtitles, the template's.
    assert "Traçabilité obligatoire - SharePoint / Confluence" in texts(slides[-1])
    assert (
        "Interactions entre 3 applications et briques du SI - flux, interfaces et données "
        "échangées" in texts(slides[3])
    )


def test_without_the_application_view_the_contents_has_six_sections() -> None:
    slides = deck_slides(deck(app_view=False))
    assert not any(title(s).startswith("Vue applicative") for s in slides)
    contents = texts(slides[0])
    assert contents[-1] == "6 - Sources consultées et limites"


def test_nothing_leaves_the_slide() -> None:
    for slide in deck_slides(deck(lines=14, contacts=60)):
        for item in slide.items:
            if isinstance(item, Box):
                box = (item.x, item.y, item.x + item.w, item.y + item.h)
            elif isinstance(item, Stroke):
                box = (
                    min(item.x1, item.x2),
                    min(item.y1, item.y2),
                    max(item.x1, item.x2),
                    max(item.y1, item.y2),
                )
            else:
                box = (item.x, item.y, item.x + sum(item.widths), item.y + sum(item.heights))
            assert box[0] >= 0 and box[1] >= 0, item
            assert box[2] <= WIDTH + 1e-6 and box[3] <= HEIGHT + 1e-6, item


def test_what_the_sources_lack_is_orange() -> None:
    slides = deck_slides(deck())
    process = slides[4]
    boxes = [i for i in process.items if isinstance(i, Box)]
    header = next(b for b in boxes if b.text == "ORDONNANCEUR ?")
    assert header.fill == ORANGE
    missing = next(b for b in boxes if b.text == "planifie ? (à confirmer)")
    assert (missing.fill, missing.color) == (ORANGE, "FFFFFF")
    step = next(b for b in boxes if b.text == "Notifie le client")
    assert step.color == ORANGE
    # Its shape, under its words: dashed, orange.
    shape = next(
        b
        for b in boxes
        if b.line == ORANGE and b.x <= step.x <= b.x + b.w and b.y <= step.y <= b.y + b.h
    )
    assert shape.dashed


def test_a_process_too_long_for_a_slide_is_split_its_business_layer_repeated() -> None:
    slides = deck_slides(deck(lines=14))
    parts = [title(s) for s in slides if "process A" in title(s)]
    assert parts == [
        "Parcours Paiement CB - Vue globale - process A : A - INITIALISATION DU PAIEMENT (1/2)",
        "Parcours Paiement CB - Vue globale - process A : A - INITIALISATION DU PAIEMENT (2/2)",
    ]
    for slide in slides:
        if "process A" in title(slide):
            assert "Choisit de payer par Carte Bancaire" in texts(slide)


def test_on_a_process_slide_lifelines_are_behind_and_a_red_line_leads_to_the_front() -> None:
    process = deck_slides(deck())[4]
    order = process.items
    lifelines = [n for n, i in enumerate(order) if isinstance(i, Stroke) and i.color == LIFELINE]
    services = [n for n, i in enumerate(order) if isinstance(i, Box) and i.fill == SERVICE]
    banners = [n for n, i in enumerate(order) if isinstance(i, Box) and i.text.startswith("1 - ")]
    assert lifelines and services and banners
    assert max(lifelines) < min(banners) < min(services)
    assert any(isinstance(i, Box) and i.line == RED and i.line_width == 2 for i in order)
    reds = [i for i in order if isinstance(i, Stroke) and i.color == RED]
    assert len(reds) == 3 and reds[-1].arrow
    # Arrows between service boxes are horizontal: one line, one chain.
    assert all(i.y1 == i.y2 for i in order if isinstance(i, Stroke) and i.arrow and i.color != RED)


def test_no_band_title_is_struck_through_by_the_front_s_line() -> None:
    # 2026-10-25: a band of one line, another of two — the front's middle
    # fell on « 2 - PAIEMENT », and its line struck the title through.
    content = deck().model_dump()
    first = content["processes"][0]
    chain = first["bands"][0]["lines"][0]
    first["bands"] = [
        {"title": "Vérification du solde", "lines": [chain]},
        {"title": "Paiement", "lines": [chain, chain]},
    ]
    process = next(
        s for s in deck_slides(ArchitectureDeck.model_validate(content)) if "process A" in title(s)
    )
    banners = [i for i in process.items if isinstance(i, Box) and i.text[:4] in ("1 - ", "2 - ")]
    assert len(banners) == 2
    for stroke in (i for i in process.items if isinstance(i, Stroke) and i.y1 == i.y2):
        for banner in banners:
            across = (
                min(stroke.x1, stroke.x2) < banner.x + banner.w
                and max(stroke.x1, stroke.x2) > banner.x
            )
            assert not (across and banner.y <= stroke.y1 <= banner.y + banner.h), banner.text


def test_a_step_s_words_never_run_under_its_arrow() -> None:
    # 2026-10-25: four lines in a narrow step ran under its « ⇒ ».
    content = deck().model_dump()
    long = "Vérification d'éligibilité du signalement selon le modèle de menace du projet"
    for grouping in content["business"]["groupings"]:
        grouping["steps"] = [{"label": long} for _ in range(4)]
    slides = deck_slides(ArchitectureDeck.model_validate(content))
    arrows = [
        i
        for s in slides
        for i in s.items
        if isinstance(i, Stroke) and i.arrow and i.width == 0.5 and i.y1 == i.y2
    ]
    assert arrows
    for slide in slides:
        for arrow in (
            i for i in slide.items if isinstance(i, Stroke) and i.arrow and i.width == 0.5
        ):
            for box in slide.items:
                if not isinstance(box, Box) or box.text != long:
                    continue
                if box.x < arrow.x2 and box.x + box.w > arrow.x1 and box.y < arrow.y1 + 0.5:
                    assert box.y > arrow.y1, "a step's words start above its arrow"


def test_the_raci_fills_its_slide_before_the_next() -> None:
    slides = deck_slides(deck(contacts=60))
    raci = [s for s in slides if title(s).startswith("Contacts et responsabilités")]
    assert len(raci) >= 2 and title(raci[1]).endswith("(suite)")
    first = next(i for i in raci[0].items if isinstance(i, Table))
    assert first.y + sum(first.heights) > HEIGHT - 0.7
    # The note, under the last one only.
    assert any("Contacts à confirmer" in t for t in texts(raci[-1]))
    assert not any("Contacts à confirmer" in t for t in texts(raci[0]))


# ── the files ───────────────────────────────────────────────────────────────


def test_the_powerpoint_is_native_and_the_pdf_has_a_page_per_slide() -> None:
    drawn = render(deck())
    assert drawn.slides == len(deck_slides(deck()))
    presentation = Presentation(BytesIO(drawn.pptx))
    assert len(presentation.slides) == drawn.slides
    assert round(presentation.slide_width / 914400, 3) == 13.333
    kinds = {shape.shape_type for slide in presentation.slides for shape in slide.shapes}
    # Shapes, connectors and tables — never a picture.
    assert 13 not in kinds
    said = [
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text
    ]
    assert "ORDONNANCEUR ?" in said
    tables = [s for slide in presentation.slides for s in slide.shapes if s.has_table]
    assert tables and tables[0].table.cell(0, 0).text == "N"
    assert drawn.pdf.startswith(b"%PDF")
    assert drawn.pdf.count(b"/Type /Page\n") + drawn.pdf.count(b"/Type /Page ") >= drawn.slides


def test_the_pdf_says_what_helvetica_cannot_draw_in_what_it_can() -> None:
    assert plain("Front → NEXUS ⇒ FACTURATION") == "Front -> NEXUS => FACTURATION"
    assert plain("Règle « À confirmer » — œuvre") == "Règle « À confirmer » — œuvre"
    assert plain("✓ 漢") == "v ?"


def test_no_em_dash_is_drawn_whatever_the_model_wrote() -> None:
    # 2026-10-26: « — » reads as a machine's on a client's slide.
    content = deck().model_dump()
    content["parcours"] = "Paiement CB — Digital"
    content["business"]["trigger"] = "Le client — ou un conseiller — paie"
    content["rules"][0]["regle"] = "Statut : en cours — à vérifier"
    drawn = ArchitectureDeck.model_validate(content)
    assert drawn.parcours == "Paiement CB : Digital"
    assert drawn.business.trigger == "Le client (ou un conseiller) paie"
    assert drawn.rules[0].regle == "Statut : en cours, à vérifier"
    assert "—" not in json.dumps(drawn.model_dump(), ensure_ascii=False)
