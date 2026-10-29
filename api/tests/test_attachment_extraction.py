"""What a file joined to an intention says, as text a model reads
(2026-10-28, Haitam: « que le LLM et LEKA lisent bien le fichier, n'importe
lequel »): each kind read for real — a PDF drawn here, Word, Excel and
PowerPoint files made here, an image and a scanned page read by OCR where
Tesseract is installed (the server's image has it), said unreadable where
it is not, never read as empty.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from PIL import Image, ImageDraw, ImageFont

from leka.platform import DOCX, PPTX, XLSX, extraction, ocr_ready, text_of

OCR = pytest.mark.skipif(not ocr_ready(), reason="Tesseract is not on this machine")


def _pdf(lines: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    page = canvas.Canvas(out, pagesize=A4)
    for number, line in enumerate(lines):
        page.drawString(72, 760 - 20 * number, line)
    page.showPage()
    page.save()
    return out.getvalue()


def _image(text: str) -> bytes:
    picture = Image.new("RGB", (1400, 220), "white")
    pen = ImageDraw.Draw(picture)
    font: ImageFont.FreeTypeFont | ImageFont.ImageFont
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 56)
    except OSError:
        font = ImageFont.load_default(size=56)
    pen.text((40, 70), text, fill="black", font=font)
    out = io.BytesIO()
    picture.save(out, format="PNG")
    return out.getvalue()


def _scanned_pdf(text: str) -> bytes:
    """A page that is only a picture of its text, as a scanner makes it."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    page = canvas.Canvas(out, pagesize=(1400, 220))
    page.drawImage(ImageReader(io.BytesIO(_image(text))), 0, 0, 1400, 220)
    page.showPage()
    page.save()
    return out.getvalue()


def _zip(parts: dict[str, str]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return out.getvalue()


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
S = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def test_a_text_a_csv_and_json_are_read_as_written() -> None:
    assert text_of("text/plain", b"Le solde est faux.\n").text == "Le solde est faux."
    csv = text_of("text/csv", b"service,p99\npayment-api,2100\n")
    assert (csv.how, csv.text) == ("text", "service,p99\npayment-api,2100")
    said = text_of("application/json", b'{"service": "payment-api", "p99": 2100}')
    assert '"service": "payment-api"' in said.text


def test_a_pdf_is_read_page_by_page() -> None:
    read = text_of("application/pdf", _pdf(["Contrat de service", "SLA : 99,95 %"]))
    assert read.how == "pdf"
    assert "[Page 1]" in read.text
    assert "Contrat de service" in read.text and "99,95" in read.text


def test_word_excel_and_powerpoint_are_read_by_their_text() -> None:
    document = _zip(
        {
            "word/document.xml": (
                f"<w:document {W}><w:body>"
                "<w:p><w:r><w:t>Plan de migration</w:t></w:r></w:p>"
                "<w:p><w:r><w:t>Étape 1 : </w:t></w:r>"
                "<w:r><w:t>geler postgres-prod</w:t></w:r></w:p>"
                "</w:body></w:document>"
            )
        }
    )
    word = text_of(DOCX, document)
    assert (word.how, word.text) == ("office", "Plan de migration\nÉtape 1 : geler postgres-prod")

    sheet = _zip(
        {
            "xl/sharedStrings.xml": (
                f"<sst {S}><si><t>service</t></si><si><t></t></si><si><t>payment-api</t></si></sst>"
            ),
            "xl/worksheets/sheet1.xml": (
                f"<worksheet {S}><sheetData>"
                '<row r="1"><c t="s"><v>0</v></c><c><v>99.95</v></c></row>'
                '<row r="2"><c t="s"><v>2</v></c><c t="inlineStr"><is><t>dégradé</t></is></c></row>'
                "</sheetData></worksheet>"
            ),
        }
    )
    excel = text_of(XLSX, sheet)
    # A shared string left empty keeps its place: « payment-api » is the 3rd.
    assert excel.text == "[sheet1]\nservice | 99.95\npayment-api | dégradé"

    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Architecture cible"
    slide.placeholders[1].text = "payment-api dépend de postgres-prod"
    out = io.BytesIO()
    deck.save(out)
    slides = text_of(PPTX, out.getvalue())
    assert slides.text.startswith("[Diapositive 1]\nArchitecture cible")
    assert "payment-api dépend de postgres-prod" in slides.text


def test_a_file_that_cannot_be_read_is_said_so_never_read_as_empty() -> None:
    broken = text_of("application/pdf", b"%PDF-1.7 nothing after")
    assert (broken.how, broken.text) == ("unreadable", "")
    assert text_of("application/zip", b"PK").why == "type"
    # An Office file that unpacks to more than it may.
    bomb = _zip({"word/document.xml": f"<w:document {W}>" + " " * 2000 + "</w:document>"})
    previous = extraction.MOST_PART
    extraction.MOST_PART = 1000
    try:
        assert text_of(DOCX, bomb).why == "error"
    finally:
        extraction.MOST_PART = previous


def test_a_long_file_is_cut_and_said_cut() -> None:
    read = text_of("text/plain", ("ligne\n" * 20_000).encode())
    assert read.cut and len(read.text) == extraction.MOST_CHARS


@pytest.mark.skipif(ocr_ready(), reason="Tesseract is on this machine: it reads")
def test_an_image_without_ocr_here_is_said_unreadable() -> None:
    assert text_of("image/png", _image("Facture 2026")).why == "no_ocr"
    assert text_of("application/pdf", _scanned_pdf("Facture 2026")).why == "scanned_no_ocr"


@OCR
def test_an_image_is_read_by_ocr() -> None:
    read = text_of("image/png", _image("Le service payment-api est lent"))
    assert read.how == "ocr"
    assert "payment-api" in read.text and "lent" in read.text


@OCR
def test_a_scanned_page_is_read_by_ocr() -> None:
    read = text_of("application/pdf", _scanned_pdf("Incident du 12 octobre"))
    assert read.how == "pdf+ocr"
    assert "Incident" in read.text and "octobre" in read.text
